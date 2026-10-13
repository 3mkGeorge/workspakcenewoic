"""Document export generation service for StorePro.

Extracts OpenPyXL (Excel) and ReportLab (PDF) financial and inventory
report document building out of HTTP routes into reusable service functions.
"""

from io import BytesIO
from datetime import datetime, timezone
from zoneinfo import ZoneInfo
from flask import current_app
from services.analytics import report_data
from services.i18n import translate

REPORTS = {
    "sales": "Sales Report",
    "purchases": "Purchases Report",
    "inventory": "Inventory Report",
    "products": "Product Report",
    "customers": "Customer Report",
    "suppliers": "Supplier Report",
    "profit": "Profit and Loss",
}


def generate_report_xlsx(kind, start, end, settings):
    """Generate an Excel (.xlsx) workbook for the specified report kind and date range.

    Returns:
        (BytesIO_stream, filename)
    """
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    headers, rows, money_cols, _ = report_data(kind, start, end)
    headers = [translate(h) for h in headers]
    if kind == "profit":
        rows = [[translate(r[0]), r[1]] for r in rows]

    filename = f"StorePro-{kind}-{start}-{end}.xlsx"
    book = Workbook()
    sheet = book.active
    sheet.title = kind.title()
    sheet.sheet_view.rightToLeft = settings.get("language") == "ar"

    sheet.append([
        settings.get("store_name", "StorePro"),
        translate(REPORTS.get(kind, "Report")),
        str(start),
        str(end),
        settings.get("currency", "EGP"),
    ])
    sheet.append(headers)

    for data in rows:
        row = [v / 100 if i in money_cols else v for i, v in enumerate(data)]
        for i, v in enumerate(row):
            if isinstance(v, str) and v.startswith(("=", "+", "-", "@")):
                row[i] = "'" + v
        sheet.append(row)

    for cell in sheet[2]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1266ED")

    for col in range(1, len(headers) + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 25

    for row in sheet.iter_rows(min_row=3):
        for i, cell in enumerate(row):
            if i in money_cols:
                cell.number_format = '"' + settings.get("currency", "EGP") + '" #,##0.00'
            elif isinstance(cell.value, datetime):
                cell.value = (
                    cell.value.replace(tzinfo=timezone.utc)
                    .astimezone(ZoneInfo("Africa/Cairo"))
                    .replace(tzinfo=None)
                )
                cell.number_format = "dd/mm/yyyy hh:mm"

    for row in sheet.iter_rows():
        for cell in row:
            if isinstance(cell.value, str) and cell.value.startswith(("=", "+", "-", "@")):
                cell.value = "'" + cell.value

    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{sheet.max_row}"

    stream = BytesIO()
    book.save(stream)
    stream.seek(0)
    return stream, filename


def generate_report_pdf(kind, start, end, settings):
    """Generate a PDF document for the specified report kind and date range.

    Returns:
        (BytesIO_stream, filename)
    """
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
    )
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from xml.sax.saxutils import escape
    from bidi.algorithm import get_display
    import arabic_reshaper

    headers, rows, money_cols, _ = report_data(kind, start, end)
    headers = [translate(h) for h in headers]
    if kind == "profit":
        rows = [[translate(r[0]), r[1]] for r in rows]

    filename = f"StorePro-{kind}-{start}-{end}.pdf"
    font = current_app.config["ROOT"] / "static/vendor/NotoSansArabic.ttf"
    pdfmetrics.registerFont(TTFont("StorePro", str(font)))
    styles = getSampleStyleSheet()
    for name in ("Normal", "Title"):
        styles[name].fontName = "StorePro"
    styles["Normal"].fontSize = 8
    styles["Normal"].leading = 12

    def cell(v):
        s = (
            current_app.jinja_env.filters["date"](v, True)
            if isinstance(v, datetime)
            else str(v)
        )
        if any("\u0600" <= c <= "\u06ff" for c in s):
            s = get_display(arabic_reshaper.reshape(s))
        return Paragraph(escape(s), styles["Normal"])

    data = [[cell(v) for v in headers]] + [
        [
            cell(
                f"{settings.get('currency', 'EGP')} {v / 100:,.2f}" if i in money_cols else v
            )
            for i, v in enumerate(row)
        ]
        for row in rows
    ]

    stream = BytesIO()
    doc = SimpleDocTemplate(
        stream,
        pagesize=landscape(A4),
        rightMargin=25,
        leftMargin=25,
        topMargin=30,
        bottomMargin=30,
    )
    table = Table(
        data,
        colWidths=[(landscape(A4)[0] - 50) / len(headers)] * len(headers),
        repeatRows=1,
    )
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eaf1ff")),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#f6f8fc")],
                ),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#1266ed")),
            ]
        )
    )

    def footer(canvas, document):
        canvas.setFont("Helvetica", 8)
        canvas.drawRightString(815, 15, f"StorePro | {document.page}")

    doc.build(
        [
            cell(settings.get("store_name", "StorePro") + " — " + translate(REPORTS.get(kind, "Report"))),
            cell(f"{start:%d/%m/%Y} – {end:%d/%m/%Y} | {settings.get('currency', 'EGP')}"),
            Spacer(1, 18),
            table,
        ],
        onFirstPage=footer,
        onLaterPages=footer,
    )
    stream.seek(0)
    return stream, filename
