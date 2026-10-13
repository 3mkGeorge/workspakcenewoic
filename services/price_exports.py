from io import BytesIO
from pathlib import Path
from flask import current_app
from services.prices import date_text, local_time, money


def export_history(product, rows, fmt, currency):
    stream = BytesIO()
    if fmt == "xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter

        book = Workbook()
        sheet = book.active
        sheet.title = "Price History"
        sheet.append(
            [
                "StorePro · Price History",
                product.name,
                "Product ID",
                product.id,
                "SKU",
                product.sku,
            ]
        )
        sheet.append(
            [
                "History ID",
                "Product ID",
                "Product Name",
                "Date & Time (Cairo)",
                "Old Cost Price",
                "New Cost Price",
                "Old Selling Price",
                "New Selling Price",
                "Selling Difference",
                "Selling Change %",
                "Cost Difference",
                "Cost Change %",
                "Changed By",
                "User ID",
                "Currency",
                "Source",
            ]
        )
        for r in rows:
            sheet.append(
                [
                    r.id,
                    r.product_id,
                    r.product_name,
                    local_time(r.created_at).replace(tzinfo=None),
                    r.old_cost / 100,
                    r.new_cost / 100,
                    r.old_price / 100,
                    r.new_price / 100,
                    r.change_amount / 100,
                    float(r.change_percentage) / 100
                    if r.change_percentage is not None
                    else "N/A (previous price zero)",
                    r.cost_change_amount / 100,
                    float(r.cost_change_percentage) / 100
                    if r.cost_change_percentage is not None
                    else "N/A (previous cost zero)",
                    r.changed_by,
                    r.user_id,
                    r.currency,
                    r.source,
                ]
            )
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"  # Never interpret names as formulas.
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for cell in sheet[2]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1266ED")
        for row in sheet.iter_rows(min_row=3):
            row[3].number_format = "dd/mm/yyyy hh:mm"
            for index in (4, 5, 6, 7, 8, 10):
                row[index].number_format = '#,##0.00 "' + str(row[14].value) + '"'
            for index in (9, 11):
                row[index].number_format = "0.00%"
        for col in range(1, 17):
            sheet.column_dimensions[get_column_letter(col)].width = (
                24 if col not in (1, 2, 14) else 12
            )
        sheet.freeze_panes = "E3"
        sheet.auto_filter.ref = f"A2:P{max(2, sheet.max_row)}"
        sheet.print_title_rows = "1:2"
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        book.save(stream)
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        from reportlab.platypus import (
            SimpleDocTemplate,
            Paragraph,
            Spacer,
            Table,
            TableStyle,
            KeepTogether,
        )
        from reportlab.lib.pagesizes import landscape, A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib import colors
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from xml.sax.saxutils import escape
        import arabic_reshaper
        from bidi.algorithm import get_display

        font = Path(current_app.root_path) / "static/vendor/NotoSansArabic.ttf"
        if "StoreProHistory" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("StoreProHistory", str(font)))
        styles = getSampleStyleSheet()
        for name in ("Normal", "Title", "Heading2"):
            styles[name].fontName = "StoreProHistory"
        styles["Normal"].fontSize = 8
        styles["Normal"].leading = 13
        styles["Heading2"].fontSize = 10
        styles["Heading2"].leading = 16

        def cell(value, style="Normal"):
            text = str(value)
            if any("\u0600" <= c <= "\u06ff" for c in text):
                text = get_display(arabic_reshaper.reshape(text))
            return Paragraph(escape(text), styles[style])

        story = [
            cell("Price History", "Title"),
            cell(
                f"{product.name} · SKU {product.sku} · Product #{product.id}",
                "Heading2",
            ),
            cell(f"{len(rows)} recorded changes · Africa/Cairo · DD/MM/YYYY HH:mm"),
            Spacer(1, 15),
        ]
        if not rows:
            story.append(
                cell(
                    "No price changes recorded. Tracking begins when the first price change is saved."
                )
            )
        for r in rows:

            def pct(value):
                return (
                    f"{value:+.2f}%"
                    if value is not None
                    else "N/A · previous value zero"
                )

            data = [
                ["Price", "Old price", "New price", "Difference", "Percentage change"],
                [
                    "Selling",
                    money(r.old_price, r.currency),
                    money(r.new_price, r.currency),
                    ("+" if r.change_amount > 0 else "")
                    + money(r.change_amount, r.currency),
                    pct(r.change_percentage),
                ],
                [
                    "Cost",
                    money(r.old_cost, r.currency),
                    money(r.new_cost, r.currency),
                    ("+" if r.cost_change_amount > 0 else "")
                    + money(r.cost_change_amount, r.currency),
                    pct(r.cost_change_percentage),
                ],
            ]
            table = Table(
                [[cell(v) for v in line] for line in data],
                colWidths=[65, 165, 165, 165, 170],
            )
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#e7effb")),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#d6e0ed")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 7),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 7),
                    ]
                )
            )
            story.append(
                KeepTogether(
                    [
                        cell(
                            f"#{r.id} · {date_text(r.created_at)} · {r.product_name}",
                            "Heading2",
                        ),
                        cell(
                            f"Changed by: {r.changed_by} · User ID: {r.user_id if r.user_id is not None else 'System'} · {r.source}"
                        ),
                        Spacer(1, 5),
                        table,
                        Spacer(1, 18),
                    ]
                )
            )

        def footer(canvas, doc):
            canvas.setFont("Helvetica", 8)
            canvas.setFillColor(colors.HexColor("#65758d"))
            canvas.drawString(35, 18, "StorePro · Append-only price history")
            canvas.drawRightString(807, 18, f"Page {doc.page}")

        SimpleDocTemplate(
            stream,
            pagesize=landscape(A4),
            rightMargin=35,
            leftMargin=35,
            topMargin=30,
            bottomMargin=35,
            title=f"Price History · {product.name}",
            author="StorePro",
        ).build(story, onFirstPage=footer, onLaterPages=footer)
        mime = "application/pdf"
    stream.seek(0)
    return stream, mime
